import os
import sys
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

OUT_DIR = '01_每週教學簡報_PPT'
os.makedirs(OUT_DIR, exist_ok=True)

COLOR_NAVY = RGBColor(30, 58, 138)
COLOR_AMBER = RGBColor(217, 119, 6)
COLOR_SLATE = RGBColor(15, 23, 42)
COLOR_MUTED = RGBColor(100, 116, 139)
COLOR_BG_CARD = RGBColor(248, 250, 252)
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_BORDER = RGBColor(203, 213, 225)
COLOR_EMERALD = RGBColor(5, 150, 105)
COLOR_PURPLE = RGBColor(124, 58, 237)

FONT_FAMILY = '微軟正黑體'

def load_antigravity_missions():
    missions_path = os.path.join('data', 'antigravity_missions.json')
    if os.path.exists(missions_path):
        with open(missions_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
            return {m['week']: m for m in data}
    return {}

def create_deck(week_num, week_title, slides_data, mission=None):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    deck_slides = []
    n_original = len(slides_data)
    
    for idx, slide_info in enumerate(slides_data):
        if mission and idx == (n_original - 1):
            deck_slides.append({'type': 'antigravity', 'mission': mission})
        deck_slides.append(slide_info)
        
    if mission and not any(s.get('type') == 'antigravity' for s in deck_slides):
        deck_slides.append({'type': 'antigravity', 'mission': mission})
        
    for s_idx, slide_info in enumerate(deck_slides):
        stype = slide_info.get('type', 'content')
        if stype == 'cover':
            create_cover_slide(prs, week_num, slide_info)
        elif stype == 'info_cards':
            create_cards_slide(prs, week_num, slide_info)
        elif stype == 'antigravity':
            create_antigravity_slide(prs, week_num, slide_info['mission'])
        else:
            create_standard_content_slide(prs, week_num, slide_info)
            
    clean_title = week_title.replace(' ', '_').replace('/', '_')
    filename = f'ERP_第{week_num:02d}週_{clean_title}.pptx'
    filepath = os.path.join(OUT_DIR, filename)
    prs.save(filepath)
    print(f'[OK] 已生成第 {week_num:02d} 週簡報 ({len(deck_slides)} 頁): {filepath}')
    return filepath

def create_cover_slide(prs, week_num, info):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_NAVY
    bg.line.fill.background()
    
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(0.18), Inches(5.0))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_AMBER
    bar.line.fill.background()
    
    tx_box = slide.shapes.add_textbox(Inches(1.3), Inches(1.1), Inches(11.2), Inches(5.2))
    tf = tx_box.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.space_after = Pt(8)
    r0 = p0.add_run()
    r0.text = '萬能科技大學 企業管理系 ｜ 11501 企業資源規劃 (ERP)'
    r0.font.name = FONT_FAMILY
    r0.font.size = Pt(28)
    r0.font.bold = True
    r0.font.color.rgb = COLOR_AMBER
    
    p1 = tf.add_paragraph()
    p1.space_after = Pt(12)
    r1 = p1.add_run()
    r1.text = f'第 {week_num:02d} 週：' + info.get('title', '')
    r1.font.name = FONT_FAMILY
    r1.font.size = Pt(40)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_WHITE
    
    p2 = tf.add_paragraph()
    p2.space_after = Pt(20)
    r2 = p2.add_run()
    r2.text = info.get('subtitle', '企業流程整合 ✕ Agentic AI 實務應用')
    r2.font.name = FONT_FAMILY
    r2.font.size = Pt(30)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(226, 232, 240)
    
    p3 = tf.add_paragraph()
    p3.space_after = Pt(10)
    r3 = p3.add_run()
    r3.text = '班級：進企四系4甲 ｜ 時間：週四 16:20 ~ 17:50 (電腦教室)'
    r3.font.name = FONT_FAMILY
    r3.font.size = Pt(28)
    r3.font.color.rgb = RGBColor(203, 213, 225)
    
    p4 = tf.add_paragraph()
    r4 = p4.add_run()
    r4.text = '授課教師：邱俊維 博士 ｜ ★ 考取 AI 賦能 ERP 或相關證照直接加分！'
    r4.font.name = FONT_FAMILY
    r4.font.size = Pt(28)
    r4.font.bold = True
    r4.font.color.rgb = COLOR_AMBER

def add_header(slide, badge_text, title_text, week_num):
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_AMBER
    top_bar.line.fill.background()
    
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(2.6), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = COLOR_NAVY
    badge.line.fill.background()
    tf_b = badge.text_frame
    tf_b.word_wrap = True
    p_b = tf_b.paragraphs[0]
    p_b.alignment = PP_ALIGN.CENTER
    r_b = p_b.add_run()
    r_b.text = badge_text
    r_b.font.name = FONT_FAMILY
    r_b.font.size = Pt(20)
    r_b.font.bold = True
    r_b.font.color.rgb = COLOR_WHITE
    
    tx_r = slide.shapes.add_textbox(Inches(7.5), Inches(0.3), Inches(5.0), Inches(0.5))
    tf_r = tx_r.text_frame
    p_r = tf_r.paragraphs[0]
    p_r.alignment = PP_ALIGN.RIGHT
    r_r = p_r.add_run()
    r_r.text = f'進企四系4甲 ｜ 第 {week_num:02d} 週'
    r_r.font.name = FONT_FAMILY
    r_r.font.size = Pt(20)
    r_r.font.color.rgb = COLOR_MUTED
    
    tx_t = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
    tf_t = tx_t.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    r_t = p_t.add_run()
    r_t.text = title_text
    r_t.font.name = FONT_FAMILY
    r_t.font.size = Pt(34)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_NAVY

def create_standard_content_slide(prs, week_num, info):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, info.get('badge', '核心精講'), info.get('title', ''), week_num)
    
    bullets = info.get('bullets', [])
    footer_note = info.get('footer')
    
    n_bullets = len(bullets)
    if n_bullets <= 4:
        sp_after = Pt(20)
        sp_before = Pt(4)
    elif n_bullets <= 5:
        sp_after = Pt(14)
        sp_before = Pt(2)
    else:
        sp_after = Pt(8)
        sp_before = Pt(1)
        
    content_h = Inches(4.3) if footer_note else Inches(5.3)
    tx = slide.shapes.add_textbox(Inches(0.8), Inches(1.75), Inches(11.733), content_h)
    tf = tx.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    
    for idx, b_text in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = sp_after
        p.space_before = sp_before
        
        r_bullet = p.add_run()
        r_bullet.text = '• '
        r_bullet.font.name = FONT_FAMILY
        r_bullet.font.size = Pt(28)
        r_bullet.font.bold = True
        r_bullet.font.color.rgb = COLOR_AMBER
        
        r_content = p.add_run()
        r_content.text = b_text
        r_content.font.name = FONT_FAMILY
        r_content.font.size = Pt(28)
        r_content.font.color.rgb = COLOR_SLATE
        
    if footer_note:
        f_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.15), Inches(11.733), Inches(1.05))
        f_box.fill.solid()
        f_box.fill.fore_color.rgb = RGBColor(254, 243, 199)
        f_box.line.color.rgb = COLOR_AMBER
        f_box.line.width = Pt(1.5)
        tf_f = f_box.text_frame
        tf_f.word_wrap = True
        tf_f.margin_top = Inches(0.08)
        tf_f.margin_bottom = Inches(0.08)
        tf_f.margin_left = Inches(0.18)
        tf_f.margin_right = Inches(0.18)
        p_f = tf_f.paragraphs[0]
        p_f.alignment = PP_ALIGN.LEFT
        r_f = p_f.add_run()
        r_f.text = '💡 關鍵提點：' + footer_note
        r_f.font.name = FONT_FAMILY
        r_f.font.size = Pt(28)
        r_f.font.bold = True
        r_f.font.color.rgb = RGBColor(146, 64, 14)

def create_cards_slide(prs, week_num, info):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, info.get('badge', '重點對比'), info.get('title', ''), week_num)
    
    cards = info.get('cards', [])
    n_cards = len(cards)
    if n_cards == 0: return
    
    gap = 0.4
    total_w = 11.733
    card_w = (total_w - (n_cards - 1) * gap) / n_cards
    top = 1.8
    card_h = 5.0
    
    for i, card in enumerate(cards):
        left = 0.8 + i * (card_w + gap)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(card_w), Inches(card_h))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_BG_CARD
        box.line.color.rgb = COLOR_BORDER
        box.line.width = Pt(2)
        
        header_h = 0.8
        h_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(card_w), Inches(header_h))
        h_bar.fill.solid()
        c_color = COLOR_NAVY if card.get('color') == 'navy' else (COLOR_PURPLE if card.get('color') == 'purple' else COLOR_EMERALD)
        h_bar.fill.fore_color.rgb = c_color
        h_bar.line.fill.background()
        
        tf_h = h_bar.text_frame
        p_h = tf_h.paragraphs[0]
        p_h.alignment = PP_ALIGN.CENTER
        r_h = p_h.add_run()
        r_h.text = card.get('title', '')
        r_h.font.name = FONT_FAMILY
        r_h.font.size = Pt(28)
        r_h.font.bold = True
        r_h.font.color.rgb = COLOR_WHITE
        
        tx_body = slide.shapes.add_textbox(Inches(left + 0.15), Inches(top + header_h + 0.15), Inches(card_w - 0.3), Inches(card_h - header_h - 0.3))
        tf_b = tx_body.text_frame
        tf_b.word_wrap = True
        
        for b_idx, b_text in enumerate(card.get('bullets', [])):
            p = tf_b.paragraphs[0] if b_idx == 0 else tf_b.add_paragraph()
            p.space_after = Pt(14)
            p.space_before = Pt(2)
            
            r_dot = p.add_run()
            r_dot.text = '• '
            r_dot.font.name = FONT_FAMILY
            r_dot.font.size = Pt(28)
            r_dot.font.bold = True
            r_dot.font.color.rgb = c_color
            
            r_txt = p.add_run()
            r_txt.text = b_text
            r_txt.font.name = FONT_FAMILY
            r_txt.font.size = Pt(28)
            r_txt.font.color.rgb = COLOR_SLATE

def create_antigravity_slide(prs, week_num, mission):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    mission_title = mission.get('title', f'Week {week_num:02d} 實戰任務')
    add_header(slide, '實機任務', f'Antigravity 任務：{mission_title}', week_num)
    
    tx = slide.shapes.add_textbox(Inches(0.8), Inches(1.75), Inches(11.733), Inches(4.3))
    tf = tx.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    
    zip_fn = mission.get('zip_filename', '專案包')
    deliv = mission.get('deliverable', '')
    
    items = [
        ('🏢 企業情境：', mission.get('scenario', '')),
        ('🤖 AI 指引：', mission.get('prompt', '')),
        ('📦 產出交付：', f'{deliv} (存放於專案目錄)'),
        ('🎯 實作驗收：', f'解壓 {zip_fn} ➔ 執行 starter.py ➔ verify.py 驗收通過')
    ]
    
    for idx, (label, content_text) in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        p.space_before = Pt(2)
        
        r_lbl = p.add_run()
        r_lbl.text = label
        r_lbl.font.name = FONT_FAMILY
        r_lbl.font.size = Pt(28)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = COLOR_NAVY if ('🏢' in label or '🎯' in label) else (COLOR_PURPLE if '🤖' in label else COLOR_EMERALD)
        
        r_txt = p.add_run()
        if len(content_text) > 75:
            content_text = content_text[:72] + '...'
        r_txt.text = content_text
        r_txt.font.name = FONT_FAMILY
        r_txt.font.size = Pt(28)
        r_txt.font.color.rgb = COLOR_SLATE
        if '📦' in label:
            r_txt.font.bold = True
            
    f_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.15), Inches(11.733), Inches(1.05))
    f_box.fill.solid()
    f_box.fill.fore_color.rgb = RGBColor(236, 253, 245)
    f_box.line.color.rgb = COLOR_EMERALD
    f_box.line.width = Pt(1.5)
    tf_f = f_box.text_frame
    tf_f.word_wrap = True
    tf_f.margin_top = Inches(0.08)
    tf_f.margin_bottom = Inches(0.08)
    tf_f.margin_left = Inches(0.18)
    tf_f.margin_right = Inches(0.18)
    p_f = tf_f.paragraphs[0]
    p_f.alignment = PP_ALIGN.LEFT
    r_f = p_f.add_run()
    r_f.text = f'💡 電腦教室實戰：一人一機下載 {zip_fn}，動手做出真正可運行的產出！'
    r_f.font.name = FONT_FAMILY
    r_f.font.size = Pt(28)
    r_f.font.bold = True
    r_f.font.color.rgb = RGBColor(6, 95, 70)

def generate_all_weeks():
    try:
        from tools.slides_metadata import WEEKS_DATA
    except ImportError:
        import slides_metadata
        WEEKS_DATA = slides_metadata.WEEKS_DATA
        
    missions = load_antigravity_missions()
    print(f'[START] 開始自動生成 18 週教學簡報，共 {len(WEEKS_DATA)} 週，已載入 {len(missions)} 週 Antigravity 實作任務...')
    generated_files = []
    
    for w_idx, w_data in enumerate(WEEKS_DATA, 1):
        week_num = w_data['week']
        week_title = w_data['title']
        slides_list = w_data['slides']
        mission = missions.get(week_num)
        
        fpath = create_deck(week_num, week_title, slides_list, mission=mission)
        generated_files.append((week_num, week_title, fpath))
        
    print(f'\n[DONE] 全部 18 週簡報生成完畢！共生成 {len(generated_files)} 份 PPTX 檔案。')
    return generated_files

if __name__ == '__main__':
    generate_all_weeks()
