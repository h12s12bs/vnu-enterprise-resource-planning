# -*- coding: utf-8 -*-
"""
生成「期中考全部題庫(含答案)」、「期末考全部題庫(含答案)」與「全學期總題庫(含答案)」Word (.docx) 電子檔
"""
import json
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def set_run_font(run, font_name='微軟正黑體', size_pt=11, bold=False, color_rgb=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color_rgb:
        run.font.color.rgb = RGBColor(*color_rgb)

def build_exam_docx(title, subtitle, scope_desc, questions, output_paths):
    doc = Document()

    # 設定頁面邊界
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 標題
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(title)
    set_run_font(r_title, size_pt=18, bold=True, color_rgb=(30, 58, 138))

    # 副標題与資訊
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run(subtitle)
    set_run_font(r_sub, size_pt=11, bold=True, color_rgb=(71, 85, 105))

    # 考試範圍說明框
    p_desc = doc.add_paragraph()
    r_desc = p_desc.add_run(scope_desc)
    set_run_font(r_desc, size_pt=10.5, bold=False, color_rgb=(30, 41, 59))

    # 一、標準答案速查表
    p_sec1 = doc.add_paragraph()
    r_sec1 = p_sec1.add_run(f"一、標準答案速查表（共 {len(questions)} 題）")
    set_run_font(r_sec1, size_pt=13, bold=True, color_rgb=(15, 23, 42))

    # 每列顯示 10 題答案
    cols = 10
    rows = (len(questions) + cols - 1) // cols
    table = doc.add_table(rows=rows, cols=cols)
    table.style = 'Table Grid'
    for idx, q in enumerate(questions):
        r_idx = idx // cols
        c_idx = idx % cols
        cell = table.cell(r_idx, c_idx)
        p_cell = cell.paragraphs[0]
        p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_num = p_cell.add_run(f"#{idx+1:03d}: ")
        set_run_font(r_num, size_pt=9.5, bold=False, color_rgb=(100, 116, 139))
        r_ans = p_cell.add_run(f"({q['answer']})")
        set_run_font(r_ans, size_pt=10, bold=True, color_rgb=(185, 28, 28))

    doc.add_paragraph()

    # 二、完整試題、標準答案與詳解（依章節排序）
    p_sec2 = doc.add_paragraph()
    r_sec2 = p_sec2.add_run("二、各章節完整考題、標準答案與考點解析")
    set_run_font(r_sec2, size_pt=13, bold=True, color_rgb=(15, 23, 42))

    current_ch = None
    for idx, q in enumerate(questions, start=1):
        if q['chapter'] != current_ch:
            current_ch = q['chapter']
            p_ch = doc.add_paragraph()
            p_ch.paragraph_format.space_before = Pt(12)
            p_ch.paragraph_format.space_after = Pt(4)
            r_ch = p_ch.add_run(f"【第 {current_ch} 章：{q['chapter_title']}】（模組：{q.get('module', 'ERP')}）")
            set_run_font(r_ch, size_pt=13, bold=True, color_rgb=(29, 78, 216))

        # 題目與答案同行
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)

        r_qnum = p_q.add_run(f"{idx}. ")
        set_run_font(r_qnum, size_pt=11, bold=True, color_rgb=(15, 23, 42))

        r_qans = p_q.add_run(f"【答案：({q['answer']})】 ")
        set_run_font(r_qans, size_pt=11, bold=True, color_rgb=(185, 28, 28))

        r_qstem = p_q.add_run(q['question'])
        set_run_font(r_qstem, size_pt=11, bold=True, color_rgb=(15, 23, 42))

        r_qsrc = p_q.add_run(f"  [{q.get('source', '教材題庫')}]")
        set_run_font(r_qsrc, size_pt=9, bold=False, color_rgb=(100, 116, 139))

        # 四個選項
        p_opts = doc.add_paragraph()
        p_opts.paragraph_format.left_indent = Inches(0.25)
        p_opts.paragraph_format.space_before = Pt(0)
        p_opts.paragraph_format.space_after = Pt(2)

        for opt_idx, opt_text in enumerate(q['options'], start=1):
            is_correct = (opt_idx == q['answer'])
            prefix = f"({opt_idx}) "
            suffix = "   " if opt_idx < 4 else ""
            r_opt = p_opts.add_run(prefix + opt_text + suffix)
            if is_correct:
                set_run_font(r_opt, size_pt=10.5, bold=True, color_rgb=(21, 128, 61))
            else:
                set_run_font(r_opt, size_pt=10.5, bold=False, color_rgb=(51, 65, 85))

        # 詳解
        if q.get('explanation'):
            clean_exp = q['explanation'].replace('\n', ' ')
            p_exp = doc.add_paragraph()
            p_exp.paragraph_format.left_indent = Inches(0.25)
            p_exp.paragraph_format.space_before = Pt(0)
            p_exp.paragraph_format.space_after = Pt(4)
            r_exp = p_exp.add_run(f"💡 解析：{clean_exp}")
            set_run_font(r_exp, size_pt=9.5, bold=False, color_rgb=(71, 85, 105))

    for path in output_paths:
        os.makedirs(os.path.dirname(path), exist_ok=True)
        doc.save(path)
        print(f"[OK] Saved: {path}")

def main():
    with open('data/questions.json', 'r', encoding='utf-8') as f:
        all_questions = json.load(f)

    # 依章節與原始題號排序
    all_questions_sorted = sorted(all_questions, key=lambda x: (x['chapter'], x['id']))

    # 期中考範圍：第 1 章 ~ 第 8 章 (105 題)
    midterm_questions = [q for q in all_questions_sorted if q['chapter'] <= 8]

    # 期末考範圍：第 9 章 ~ 第 16 章 (95 題)
    final_questions = [q for q in all_questions_sorted if q['chapter'] >= 9]

    print(f"Midterm count (CH1-CH8): {len(midterm_questions)}")
    print(f"Final count (CH9-CH16): {len(final_questions)}")
    print(f"Total count (CH1-CH16): {len(all_questions_sorted)}")

    # 1. 期中考題庫
    build_exam_docx(
        title="萬能科技大學 11501《企業資源規劃》期中考全部題庫（含標準答案與詳解）",
        subtitle="授課教師：邱俊維 博士 ｜ 涵蓋範圍：第 1 章 ～ 第 8 章（共 105 題）",
        scope_desc=(
            "📌 命題與複習說明：\n"
            "1. 本題庫涵蓋前半學期（第 1 章至第 8 章）：資訊技術與競爭優勢、ERP 概說與 BPR、系統導入方法、系統評選、系統建置、存貨管理(MM)、銷售與收帳管理(SD)、生產規劃與控制(PP/BOM/MRP)。\n"
            "2. 期中考正式測驗將直接從本題庫 105 題中抽題命題，每題均已附上官方標準答案與考點詳解，請同學熟讀練習。\n"
            "3. 題庫來源：張緯良 著《企業資源規劃：企業e化之營運管理 4e》（前程文化官方題庫）＋ 中華企業資源規劃學會 (CERPS) 認證精選考題。"
        ),
        questions=midterm_questions,
        output_paths=[
            "static/exam_banks/ERP_Midterm_Question_Bank_With_Answers.docx",
            "03_期中期末考題庫_電子檔/萬能科大_企業資源規劃_期中考全部題庫(含標準答案與詳解_第1至8章_105題).docx"
        ]
    )

    # 2. 期末考題庫
    build_exam_docx(
        title="萬能科技大學 11501《企業資源規劃》期末考全部題庫（含標準答案與詳解）",
        subtitle="授課教師：邱俊維 博士 ｜ 涵蓋範圍：第 9 章 ～ 第 16 章（共 95 題）",
        scope_desc=(
            "📌 命題與複習說明：\n"
            "1. 本題庫涵蓋後半學期（第 9 章至第 16 章）：採購與付款管理(MM/三向比對)、在製品管理(SFC)、人力資源管理(HR)、財務會計(FI)、管理會計(CO)、企業變革管理、未來展望與新興技術、AI 時代的 ERP（Agentic AI、ReAct 與流程自動化）。\n"
            "2. 期末考正式測驗將直接從本題庫 95 題中抽題命題，每題均已附上官方標準答案與考點詳解，請同學熟讀練習。\n"
            "3. 題庫來源：張緯良 著《企業資源規劃：企業e化之營運管理 4e》（前程文化官方題庫）＋ 中華企業資源規劃學會 (CERPS) AI-ERP 認證精選考題。"
        ),
        questions=final_questions,
        output_paths=[
            "static/exam_banks/ERP_Final_Question_Bank_With_Answers.docx",
            "03_期中期末考題庫_電子檔/萬能科大_企業資源規劃_期末考全部題庫(含標準答案與詳解_第9至16章_95題).docx"
        ]
    )

    # 3. 全學期總題庫 (200題完整版)
    build_exam_docx(
        title="萬能科技大學 11501《企業資源規劃》全學期總題庫完整版（含標準答案與詳解）",
        subtitle="授課教師：邱俊維 博士 ｜ 涵蓋範圍：全冊第 1 章 ～ 第 16 章（共 200 題）",
        scope_desc=(
            "📌 命題與證照檢定說明：\n"
            "1. 本檔案完整收錄全學期 200 題單選題庫（期中考範圍第 1～8 章 105 題 ＋ 期末考範圍第 9～16 章 95 題）。\n"
            "2. 適用於期中考、期末考複習以及報考「中華企業資源規劃學會 (CERPS) ERP 規劃師 ／ AI 賦能 ERP 規劃師」專業證照檢定衝刺。\n"
            "3. 考取相關專業證照並於課程平台登錄，可享學期總成績特別加分！"
        ),
        questions=all_questions_sorted,
        output_paths=[
            "static/exam_banks/ERP_All_200_Question_Bank_With_Answers.docx",
            "03_期中期末考題庫_電子檔/萬能科大_企業資源規劃_全學期總題庫完整版(含標準答案與詳解_200題).docx"
        ]
    )

if __name__ == '__main__':
    main()
